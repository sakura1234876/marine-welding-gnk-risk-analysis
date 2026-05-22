"""
Weight Parameter Sensitivity Analysis & Clustering Visualization Script
-----------------------------------
Features:
1. Calculates the combined weight and rank of paths based on different weight parameters (a, b).
2. Computes the Kendall's Tau correlation matrix for rankings across different parameter combinations.
3. Uses the Elbow Method and K-Means to cluster the parameter combinations.
4. Applies PCA dimensionality reduction to visualize the clustering results in a 2D scatter plot.
5. Exports all analytical data and the correlation matrix (with conditional formatting) to an Excel file.
"""

import pandas as pd
import numpy as np
from scipy.stats import kendalltau
import matplotlib.pyplot as plt
import seaborn as sns
from openpyxl import load_workbook
from openpyxl.formatting.rule import ColorScaleRule
import warnings
from sklearn.cluster import KMeans
from sklearn.decomposition import PCA

# ==========================================
# 1. Configuration: File Paths & Output Settings
# ==========================================
# Input file path and sheet name
INPUT_FILE_PATH = r"<your_input_directory_here>\input_data.xlsx"
SHEET_NAME = "<your_sheet_name_here>"

# Output file paths (can be configured to a specific output folder)
OUTPUT_EXCEL = r"<your_output_directory_here>\sensitivity_analysis_results.xlsx"
OUTPUT_HEATMAP = r"<your_output_directory_here>\1_kendall_heatmap.png"
OUTPUT_ELBOW = r"<your_output_directory_here>\2_kmeans_elbow.png"
OUTPUT_SCATTER = r"<your_output_directory_here>\3_kmeans_pca_scatter.png"

# ==========================================
# 2. Plotting & Warning Settings
# ==========================================
# Using standard sans-serif for English text compatibility
plt.rcParams["font.family"] = ["sans-serif"]
plt.rcParams["axes.unicode_minus"] = False
warnings.filterwarnings("ignore", category=UserWarning, module="matplotlib")

# ==========================================
# 3. Parameter Range Settings
# ==========================================
A_START = 0.0       
A_END = 1.0        
A_STEP = 0.05       
# Use linspace to avoid floating-point precision issues and ensure unique keys
a_values = np.linspace(A_START, A_END, int((A_END - A_START) / A_STEP) + 1)
b_values = 1 - a_values

# ==========================================
# 4. Column Name Settings (Modify according to your data)
# ==========================================
PATH_ID_COL = "ID"                        # Path unique identifier column
TRANS_PROB_COL = "Transition_Probability" # Transition probability column
COUPLING_T_COL = "Coupling_Degree_T"      # Coupling degree T column

# ==========================================
# 5. Data Loading & Validation
# ==========================================
df = pd.read_excel(INPUT_FILE_PATH, sheet_name=SHEET_NAME)
# Take the first 15 rows as an example, modify or remove as needed
df = df.head(15)

required_cols = [PATH_ID_COL, TRANS_PROB_COL, COUPLING_T_COL]
missing_cols = [col for col in required_cols if col not in df.columns]
if missing_cols:
    raise ValueError(f"Missing required columns in Excel: {missing_cols}. Please check column names.")

df = df.dropna(subset=required_cols).reset_index(drop=True)
df[TRANS_PROB_COL] = pd.to_numeric(df[TRANS_PROB_COL], errors="coerce")
df[COUPLING_T_COL] = pd.to_numeric(df[COUPLING_T_COL], errors="coerce")
df = df.dropna(subset=required_cols).reset_index(drop=True)

print(f"[*] Data loaded successfully. Valid paths: {len(df)}")

# ==========================================
# 6. Calculate Combined Weights & Rankings
# ==========================================
rank_results = {}
all_weight_cols = []

for a, b in zip(a_values, b_values):
    # Keep 2 decimal places to match the 0.05 step precision and avoid duplicate keys
    weight_col_name = f"Combined_Weight_a={a:.2f}_b={b:.2f}"
    df[weight_col_name] = a * df[TRANS_PROB_COL] + b * df[COUPLING_T_COL]
    all_weight_cols.append(weight_col_name)
    
    rank_col_name = f"Rank_a={a:.2f}_b={b:.2f}"
    df[rank_col_name] = df[weight_col_name].rank(ascending=False, method='min').astype(int)
    all_weight_cols.append(rank_col_name)
    
    # Use 2 decimal places for keys to ensure uniqueness
    rank_dict = dict(zip(df[PATH_ID_COL], df[rank_col_name]))
    rank_results[f"a={a:.2f}"] = rank_dict

# Generate path ranking summary table
rank_df = pd.DataFrame(index=df[PATH_ID_COL])
for param_name, rank_dict in rank_results.items():
    rank_df[param_name] = rank_df.index.map(rank_dict)
rank_df = rank_df.reset_index().rename(columns={"index": PATH_ID_COL})
print("[*] Combined weights and ranking columns calculated.")

# ==========================================
# 7. Calculate Kendall's Tau Correlation Matrix
# ==========================================
params = list(rank_results.keys())
n_params = len(params)
kendall_matrix = np.zeros((n_params, n_params))

for i in range(n_params):
    for j in range(n_params):
        if i == j:
            kendall_matrix[i, j] = 1.0
        else:
            rank_i = list(rank_results[params[i]].values())
            rank_j = list(rank_results[params[j]].values())
            tau, _ = kendalltau(rank_i, rank_j)
            kendall_matrix[i, j] = round(tau, 4)

kendall_df = pd.DataFrame(kendall_matrix, index=params, columns=params)

# Stability ranking
avg_corr = kendall_matrix.mean(axis=1)
stability_df = pd.DataFrame({
    "Parameter_Combination": params,
    "Weight_a": a_values,
    "Weight_b": b_values,
    "Avg_Kendall_Correlation": avg_corr
}).round(4)
stability_df = stability_df.sort_values(by="Avg_Kendall_Correlation", ascending=False).reset_index(drop=True)
stability_df["Stability_Rank"] = stability_df.index + 1

# ==========================================
# 8. Plot Correlation Heatmap
# ==========================================
plt.figure(figsize=(12, 10))
ax = sns.heatmap(
    kendall_matrix, annot=True, cmap="coolwarm", vmin=-1, vmax=1, fmt=".4f",
    xticklabels=params, yticklabels=params, cbar_kws={"label": "Kendall's Tau Correlation"}
)
ax.set_title("Kendall's Tau Correlation Heatmap of Path Rankings", fontsize=14, pad=20)
ax.set_xlabel("Weight Parameter Combination", fontsize=12)
ax.set_ylabel("Weight Parameter Combination", fontsize=12)
plt.xticks(rotation=45, ha="right")
plt.yticks(rotation=0)
plt.tight_layout()
plt.savefig(OUTPUT_HEATMAP, dpi=300, bbox_inches="tight")
plt.close()
print(f"[*] 1/3: Heatmap saved to: {OUTPUT_HEATMAP}")

# ==========================================
# 9. Elbow Method for Optimal K in KMeans
# ==========================================
print("[*] Calculating Sum of Squared Errors (SSE) for different K values...")
sse = [] 
K_range = range(1, 10) # Try 1 to 9 cluster centers

for k in K_range:
    kmeans_temp = KMeans(n_clusters=k, random_state=42, n_init=10)
    kmeans_temp.fit(kendall_matrix)
    sse.append(kmeans_temp.inertia_)

plt.figure(figsize=(10, 6))
plt.plot(K_range, sse, marker='o', linestyle='-', color='b', linewidth=2, markersize=8, zorder=3)
plt.title('Elbow Method for Optimal Number of Clusters (K)', fontsize=14, pad=15)
plt.xlabel('Number of Clusters (K)', fontsize=12)
plt.ylabel('Sum of Squared Errors (SSE / Inertia)', fontsize=12)
plt.xticks(K_range)
plt.grid(True, linestyle='--', alpha=0.6, zorder=0)
plt.tight_layout()
plt.savefig(OUTPUT_ELBOW, dpi=300, bbox_inches="tight")
plt.close()
print(f"[*] 2/3: Elbow method line chart saved to: {OUTPUT_ELBOW}")

# ==========================================
# 10. K-Means Clustering & PCA Visualization
# ==========================================
num_clusters = 4
print(f"[*] Executing K-Means clustering (Current setting K={num_clusters})...")
kmeans = KMeans(n_clusters=num_clusters, random_state=42, n_init=10)
cluster_labels = kmeans.fit_predict(kendall_matrix)

# PCA dimensionality reduction for visualization
pca = PCA(n_components=2)
reduced_features = pca.fit_transform(kendall_matrix)

plt.figure(figsize=(10, 8))
scatter = plt.scatter(reduced_features[:, 0], reduced_features[:, 1], 
                      c=cluster_labels, cmap='Set1', s=150, edgecolors='k', zorder=3)

# Add parameter labels to scatter points
for i, txt in enumerate(params):
    plt.annotate(txt, (reduced_features[i, 0], reduced_features[i, 1]), 
                 xytext=(8, 8), textcoords='offset points', fontsize=11)

plt.title(f"K-Means Clustering of Weight Parameter Combinations (K={num_clusters})", fontsize=14, pad=20)
plt.xlabel("Principal Component 1 (Primary Rank Variance)", fontsize=12)
plt.ylabel("Principal Component 2 (Secondary Rank Variance)", fontsize=12)
plt.grid(True, linestyle='--', alpha=0.6, zorder=0)

# Add cluster legend
legend1 = plt.legend(*scatter.legend_elements(), loc="best", title="Cluster Segment")
plt.gca().add_artist(legend1)

plt.tight_layout()
plt.savefig(OUTPUT_SCATTER, dpi=300, bbox_inches="tight")
plt.close()
print(f"[*] 3/3: Clustering scatter plot saved to: {OUTPUT_SCATTER}")

# ==========================================
# 11. Export Results to Excel
# ==========================================
cluster_records = []
for i in range(n_params):
    cluster_records.append({
        "Parameter_Combination": params[i],
        "Weight_a": a_values[i],
        "Assigned_Cluster": f"Cluster {cluster_labels[i] + 1}"  # Add 1 for 1-based indexing
    })
cluster_df = pd.DataFrame(cluster_records).sort_values(by="Weight_a").reset_index(drop=True)
cluster_df = cluster_df.drop(columns=["Weight_a"])
print(f"[*] Clustering complete. {n_params} parameter combinations partitioned into {num_clusters} clusters.")

print("[*] Writing results to Excel file...")
with pd.ExcelWriter(OUTPUT_EXCEL, engine="openpyxl") as writer:
    df.to_excel(writer, sheet_name="RawData_Weights_Ranks", index=False)
    rank_df.to_excel(writer, sheet_name="Path_Rankings_Summary", index=False)
    kendall_df.to_excel(writer, sheet_name="Kendall_Correlation_Matrix", index=True)
    stability_df.to_excel(writer, sheet_name="Parameter_Stability_Rank", index=False)
    cluster_df.to_excel(writer, sheet_name="KMeans_Clustering_Results", index=False)

# Add Red-White-Blue conditional formatting to the correlation matrix heatmap
wb = load_workbook(OUTPUT_EXCEL)
ws = wb["Kendall_Correlation_Matrix"]
max_row = ws.max_row
max_col = ws.max_column

color_scale_rule = ColorScaleRule(
    start_type="num", start_value=-1, start_color="F8696B",
    mid_type="num", mid_value=0, mid_color="FFFFFF",
    end_type="num", end_value=1, end_color="5A8AC6"
)
# Apply to all data cells (starting from B2)
ws.conditional_formatting.add(f"B2:{chr(ord('A') + max_col - 1)}{max_row}", color_scale_rule)

wb.save(OUTPUT_EXCEL)
print(f"[*] 🚀 All analytical pipelines executed successfully! Full results saved to: {OUTPUT_EXCEL}")